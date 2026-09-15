import io
import json
import pytest
import os
import subprocess
import sys
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import (
    Asset,
    BrowserArtifact,
    BrowserSession,
    Endpoint,
    Evidence,
    EvidenceAttachment,
    Finding,
    Identity,
    Project,
    StoredRequest,
    StoredRequestRevision,
)
from app.services.coverage_matrix import coverage_dimension_summary
from app.services.evidence_attachments import (
    render_annotated_image_bytes,
    save_attachment_annotations,
    save_finding_screenshot,
)
from app.services.finding_quality import finding_quality, report_preflight
from app.services.finding_service import create_finding
from app.services.personal_settings import save_personal_settings
from app.services.request_workspace import (
    create_stored_request,
    request_revisions,
    restore_request_revision,
    update_stored_request,
)
from app.services.txb02_report import generate_txb02_docx


def _png_bytes() -> bytes:
    image = Image.new("RGB", (320, 180), (32, 36, 42))
    out = io.BytesIO()
    image.save(out, format="PNG")
    return out.getvalue()


def test_v15_request_revisions_edit_restore_and_state_change_policy():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V15 Request Revision", scope_text="api.example.test")
        db.add(project); db.commit(); db.refresh(project)
        stored = create_stored_request(
            db, project.id, "GET profile", "GET",
            "https://api.example.test/profile?view=full&ticket=QUERY-SECRET",
            headers={"Accept": "application/json", "Authorization": "Bearer OLD-SECRET"},
            source="manual",
        )
        assert len(request_revisions(db, stored.id)) == 1
        rev1 = request_revisions(db, stored.id)[0]
        assert rev1.revision_no == 1

        update_stored_request(
            db, stored,
            name="Update profile",
            method="POST",
            url="https://api.example.test/profile?ticket=••••",
            headers_text="Content-Type: application/json\nAuthorization: Bearer NEW-SECRET",
            body='{"nickname":"qa"}',
            change_note="Changed to update request for manual review only",
        )
        assert stored.policy_class == "STATE_CHANGE"
        assert "QUERY-SECRET" in stored.url
        revisions = request_revisions(db, stored.id)
        assert [x.revision_no for x in revisions[:2]] == [2, 1]
        assert "NEW-SECRET" not in stored.headers_json
        assert "NEW-SECRET" not in stored.secret_headers_encrypted

        restore_request_revision(db, stored, rev1)
        assert stored.method == "GET"
        assert stored.policy_class == "READ_ONLY"
        assert len(request_revisions(db, stored.id)) == 3
    finally:
        db.close()


@pytest.mark.parametrize("token_header", ["X-AI-Pentest-Token", "X-SnowEdge-Token"])
def test_v15_burp_ingestion_requires_token_scope_and_never_replays(token_header):
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        save_personal_settings(
            db,
            {"ai_provider": "mock"},
            burp_ingest_token="LOCAL-BURP-TOKEN",
        )
        project = Project(name="V15 Burp", scope_text="api.example.test")
        db.add(project); db.commit(); db.refresh(project)
        pid = project.id
    finally:
        db.close()

    raw = (
        "POST /profile HTTP/1.1\r\n"
        "Host: api.example.test\r\n"
        "Authorization: Bearer BURP-SECRET\r\n"
        "Content-Type: application/json\r\n\r\n"
        '{"nickname":"qa"}'
    )
    with TestClient(app) as client:
        denied = client.post(
            f"/api/projects/{pid}/burp/send",
            headers={token_header: "wrong"},
            json={"scheme": "https", "raw_request": raw},
        )
        assert denied.status_code == 401

        accepted = client.post(
            f"/api/projects/{pid}/burp/send",
            headers={token_header: "LOCAL-BURP-TOKEN"},
            json={"scheme": "https", "name": "Burp POST profile", "raw_request": raw},
        )
        assert accepted.status_code == 200
        data = accepted.json()
        assert data["policy_class"] == "STATE_CHANGE"

        outside = client.post(
            f"/api/projects/{pid}/burp/send",
            headers={token_header: "LOCAL-BURP-TOKEN"},
            json={"scheme": "https", "raw_request": "GET / HTTP/1.1\r\nHost: outside.example.test\r\n\r\n"},
        )
        assert outside.status_code == 400

    db = SessionLocal()
    try:
        stored = db.get(StoredRequest, data["stored_request_id"])
        assert stored.source == "burp_extension"
        assert stored.policy_class == "STATE_CHANGE"
        assert db.query(StoredRequestRevision).filter(StoredRequestRevision.stored_request_id == stored.id).count() == 1
    finally:
        db.close()


def test_v15_finding_quality_gate_and_report_preflight():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V15 Quality", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(
            db=db,
            project_id=project.id,
            title="Authorization candidate",
            severity="medium",
            target="https://example.test/order/1001",
            description="Short",
            recommendation="Fix auth.",
            source="authorization_testing",
            evidence_kind="authorization_candidate_summary",
            evidence_content={"candidate": True},
        )
        quality = finding_quality(db, finding)
        assert quality["score"] < 75
        assert quality["ready"] is False
        keys = {x["key"] for x in quality["missing"]}
        assert "screenshot" in keys
        assert "request" in keys or "response" in keys

        preflight = report_preflight(db, project.id)
        assert preflight["ready"] is False
        assert preflight["warning_count"] >= 1
        assert any(x["finding_id"] == finding.id for x in preflight["issues"])
    finally:
        db.close()


def test_v15_annotated_screenshot_is_non_destructive_and_word_uses_rendered_image(tmp_path, monkeypatch):
    Base.metadata.create_all(bind=engine)
    monkeypatch.setattr(
        "app.services.evidence_attachments.settings.evidence_artifact_dir",
        str(tmp_path / "evidence"),
    )
    db = SessionLocal()
    try:
        project = Project(name="V15 Annotation", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(
            db=db,
            project_id=project.id,
            title="Screenshot annotation QA",
            severity="medium",
            target="https://example.test/",
            description="A sufficiently detailed finding description for screenshot annotation quality assurance.",
            recommendation="Apply the server-side authorization fix and retest the original request.",
            source="authorization_testing",
            evidence_kind="poc_request",
            evidence_content="GET / HTTP/1.1\nHost: example.test",
        )
        original = _png_bytes()
        attachment = save_finding_screenshot(db, finding, "Original screenshot", "image/png", original)
        original_file = Path(attachment.file_path).read_bytes()

        save_attachment_annotations(
            db, attachment,
            label="图示：敏感区域已打码并标框",
            annotation_json=json.dumps([
                {"type":"redact","x":0.10,"y":0.10,"w":0.25,"h":0.20},
                {"type":"box","x":0.45,"y":0.30,"w":0.30,"h":0.30},
                {"type":"arrow","x":0.10,"y":0.80,"w":0.45,"h":-0.25},
            ]),
            sort_order=3,
        )
        rendered = render_annotated_image_bytes(attachment)
        assert rendered != original_file
        assert Path(attachment.file_path).read_bytes() == original_file
        assert attachment.sort_order == 3

        report = generate_txb02_docx(db, project)
        doc = Document(io.BytesIO(report))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for row in t.rows for c in row.cells]
        )
        assert "图示：敏感区域已打码并标框" in text
    finally:
        db.close()


def test_v15_coverage2_dimensions_expose_personal_workflow_gaps():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V15 Coverage", scope_text="api.example.test")
        db.add(project); db.commit(); db.refresh(project)
        asset = Asset(project_id=project.id, target="api.example.test", kind="host")
        db.add(asset); db.commit(); db.refresh(asset)
        db.add(Endpoint(asset_id=asset.id, method="GET", url="https://api.example.test/profile"))
        db.add(Identity(project_id=project.id, name="User A", role="User"))
        db.commit()

        summary = coverage_dimension_summary(db, project)
        rows = {x["key"]: x for x in summary["rows"]}
        assert rows["identity"]["score"] == 50
        assert rows["authorization"]["score"] == 0
        assert rows["screenshot"]["score"] == 0
        assert summary["average"] < 100
    finally:
        db.close()


def test_v15_quality_and_preflight_ui_markers_render():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V15 UI", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(
            db=db,
            project_id=project.id,
            title="UI Quality Finding",
            severity="low",
            target="https://example.test/",
            description="A sufficiently detailed description used to render the Finding Quality Gate user interface.",
            recommendation="Add the missing control and perform a documented retest.",
            source="headers_check",
            evidence_kind="http_response_headers",
            evidence_content={"content-type":"text/html"},
        )
        pid, fid = project.id, finding.id
    finally:
        db.close()

    with TestClient(app) as client:
        for url, marker in [
            (f"/projects/{pid}/requests", "请求工作台 2.0"),
            (f"/projects/{pid}/findings", "平均证据质量"),
            (f"/findings/{fid}", "Finding Quality Gate"),
            (f"/projects/{pid}/reports", "Report Preflight"),
            (f"/projects/{pid}/coverage", "Coverage 2.0"),
        ]:
            response = client.get(url)
            assert response.status_code == 200
            assert marker in response.text


def test_v15_burp_extension_source_package_exists():
    root = Path(__file__).resolve().parents[1]
    java = root / "integrations/burp-extension/src/main/java/burp/BurpExtender.java"
    readme = root / "integrations/burp-extension/README.md"
    assert java.exists()
    assert readme.exists()
    text = java.read_text(encoding="utf-8")
    assert "Send to SnowEdge" in text
    assert "X-SnowEdge-Token" in text
    assert "/burp/send" in text


def test_v15_v14_database_migrates_to_v15_head(tmp_path):
    db_path = tmp_path / "v14.db"
    sql = (
        "CREATE TABLE projects (id INTEGER PRIMARY KEY, name VARCHAR(200), scope_text TEXT, client_name VARCHAR(240), "
        "environment VARCHAR(80), engagement_type VARCHAR(80), status VARCHAR(60), start_date VARCHAR(32), end_date VARCHAR(32), "
        "authorization_note TEXT, testers_json TEXT, template_slug VARCHAR(80), created_at DATETIME);"
        "CREATE TABLE evidence_attachments (id INTEGER PRIMARY KEY, project_id INTEGER, finding_id INTEGER, evidence_id INTEGER, "
        "attachment_type VARCHAR(60), label VARCHAR(300), file_path VARCHAR(2000), mime_type VARCHAR(100), file_sha256 VARCHAR(64), "
        "size_bytes INTEGER, created_at DATETIME);"
        "CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL);"
        "INSERT INTO alembic_version(version_num) VALUES('v1_4_personal');"
    )
    bootstrap = (
        "import sqlite3\n"
        f"db=sqlite3.connect(r'{db_path}')\n"
        f"db.executescript({sql!r})\n"
        "db.commit();db.close()\n"
    )
    subprocess.run([sys.executable, "-c", bootstrap], check=True)

    env = os.environ.copy()
    env["DATABASE_URL"] = f"sqlite:///{db_path}"
    code = (
        "from sqlalchemy import inspect\n"
        "from app.db import engine\n"
        "from app.schema_migrations import ensure_schema_current\n"
        "ensure_schema_current()\n"
        "i=inspect(engine)\n"
        "tables=set(i.get_table_names())\n"
        "cols={c['name'] for c in i.get_columns('evidence_attachments')}\n"
        "assert 'stored_request_revisions' in tables\n"
        "assert {'annotation_json','sort_order'} <= cols\n"
        "with engine.connect() as c: rev=c.exec_driver_sql('SELECT version_num FROM alembic_version').scalar()\n"
        "assert rev=='v1_8_toolchain_runs'\n"
        "print(rev)\n"
    )
    result = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(Path(__file__).resolve().parents[1]),
        env=env,
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "v1_8_toolchain_runs" in result.stdout


def test_v15_browser2_runtime_metadata_renders_without_secret_values():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V15 Browser2 UI", scope_text="app.example.test")
        db.add(project); db.commit(); db.refresh(project)
        session = BrowserSession(
            project_id=project.id,
            target_url="https://app.example.test/",
            status="done",
            summary_json=json.dumps({
                "requests": 12,
                "xhr_fetch": 4,
                "route_changes": 3,
                "websockets": 1,
                "local_storage_keys": 2,
                "session_storage_keys": 1,
                "blocked_out_of_scope": 2,
            }),
        )
        db.add(session); db.commit(); db.refresh(session)
        runtime = {
            "route_timeline": [
                {"url": "https://app.example.test/", "in_scope": True},
                {"url": "https://app.example.test/dashboard?ticket=SECRET-TICKET", "in_scope": True},
            ],
            "websockets": [{
                "url": "wss://app.example.test/ws?token=SECRET-WS",
                "frames_sent": 2, "frames_received": 3,
                "bytes_sent": 88, "bytes_received": 144,
                "closed": True,
            }],
            "storage": {
                "localStorageKeys": ["qa_local_key", "theme"],
                "sessionStorageKeys": ["qa_session_key"],
            },
            "cookies": [{
                "name": "sid", "domain": "app.example.test", "path": "/",
                "expires": -1, "httpOnly": True, "secure": True, "sameSite": "Lax",
                "value_omitted": True,
            }],
            "document_security_headers": {
                "content-security-policy": "default-src 'self'",
                "x-content-type-options": "nosniff",
            },
            "dom_sha256": "a" * 64,
            "previous_dom_sha256": "b" * 64,
            "dom_similarity": 93.4,
            "secrets_persisted": False,
            "websocket_payloads_persisted": False,
        }
        db.add(BrowserArtifact(
            project_id=project.id,
            browser_session_id=session.id,
            artifact_type="runtime_metadata",
            label="Runtime metadata",
            content_text=json.dumps(runtime),
            metadata_json="{}",
        ))
        db.commit()
        pid, sid = project.id, session.id
    finally:
        db.close()

    with TestClient(app) as client:
        response = client.get(f"/projects/{pid}/browser?selected={sid}")
        assert response.status_code == 200
        assert "运行时情报" in response.text
        assert "qa_local_key" in response.text
        assert "qa_session_key" in response.text
        assert "payload omitted" in response.text
        assert "default-src" in response.text
        assert "SECRET-TICKET" not in response.text
        assert "SECRET-WS" not in response.text


def test_v15_setup_ui_saves_encrypted_burp_token():
    Base.metadata.create_all(bind=engine)
    with TestClient(app) as client:
        page = client.get("/setup")
        assert page.status_code == 200
        assert "Burp 接入令牌" in page.text
        response = client.post(
            "/setup",
            data={
                "default_template": "web-api",
                "backup_enabled": "on",
                "backup_interval_hours": "24",
                "backup_retention": "7",
                "backup_dir": "./backups",
                "browser_path": "",
                "nmap_path": "",
                "ai_provider": "mock",
                "ai_api_base": "",
                "ai_model": "",
                "ai_api_key": "",
                "burp_ingest_token": "SETUP-BURP-SECRET",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303

    db = SessionLocal()
    try:
        from app.services.personal_settings import get_personal_settings, burp_ingest_token
        prefs = get_personal_settings(db)
        assert prefs["burp_ingest_token_configured"] is True
        assert burp_ingest_token(db) == "SETUP-BURP-SECRET"
        row = db.query(__import__("app.models", fromlist=["AppPreference"]).AppPreference).filter_by(key="personal_config").one()
        assert "SETUP-BURP-SECRET" not in row.secret_encrypted
    finally:
        db.close()
