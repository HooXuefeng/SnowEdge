import asyncio
import base64
import io
import json
import os
import subprocess
import sys
from datetime import timedelta
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import (
    Asset,
    Evidence,
    EvidenceProvenance,
    ImportBatch,
    ImportRecord,
    JobWorker,
    PersistentJob,
    Project,
    Service,
    StoredRequest,
    _utcnow,
)
from app.services.evidence_chain import evidence_chain_payload
from app.services.finding_service import create_finding
from app.services.finding_states import update_finding_states
from app.services.job_engine import (
    _claim_next_job,
    enqueue_job,
    heartbeat_job,
    process_job,
    recover_orphaned_jobs,
)
from app.services.passive_imports import import_burp_xml, import_nmap_xml, import_nuclei_jsonl
from app.services.txb02_report import generate_txb02_docx


def test_v13_worker_lease_heartbeat_and_expired_lease_recovery():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        # The legacy suite shares one SQLite database; isolate queue ownership for this test.
        for stale in db.query(PersistentJob).filter(PersistentJob.status.in_(["queued", "running", "retry_wait", "cancel_requested"])).all():
            stale.status = "cancelled"
            stale.lease_expires_at = None
            stale.worker_id = ""
            stale.lease_token = ""
        db.commit()
        project = Project(name="V13 Worker Lease", scope_text="127.0.0.1")
        db.add(project); db.commit(); db.refresh(project)
        job = enqueue_job(
            db, project, "endpoint_sync",
            target="project:endpoints",
            payload={"reason": "lease-test"},
            max_attempts=1,
        )
        job_id = job.id
        project_id = project.id
    finally:
        db.close()

    claimed = asyncio.run(_claim_next_job("worker-test-A"))
    assert claimed is not None
    claimed_id, token = claimed
    assert claimed_id == job_id

    db = SessionLocal()
    try:
        project = db.get(Project, project_id)
        job = db.get(PersistentJob, job_id)
        assert job.status == "running"
        assert job.worker_id == "worker-test-A"
        assert job.lease_token == token
        assert job.lease_expires_at > _utcnow()
        worker = db.query(JobWorker).filter(JobWorker.worker_id == "worker-test-A").one()
        assert worker.status == "online"
        assert worker.active_job_id == job_id
    finally:
        db.close()

    assert heartbeat_job(job_id, "worker-test-A", token) is True
    assert heartbeat_job(job_id, "worker-test-B", token) is False

    asyncio.run(process_job(job_id, "worker-test-A", token))
    db = SessionLocal()
    try:
        job = db.get(PersistentJob, job_id)
        assert job.status == "done"
        assert job.lease_expires_at is None

        orphan = enqueue_job(
            db, project, "endpoint_sync",
            target="project:orphan",
            payload={"reason": "orphan-test"},
            max_attempts=1,
        )
        orphan.status = "running"
        orphan.worker_id = "dead-worker"
        orphan.lease_token = "dead-token"
        orphan.lease_expires_at = _utcnow() - timedelta(seconds=1)
        db.commit()
        assert recover_orphaned_jobs(db) >= 1
        db.refresh(orphan)
        assert orphan.status == "queued"
        assert orphan.worker_id == ""
        assert orphan.lease_token == ""
    finally:
        db.close()


def test_v13_evidence_integrity_and_provenance_chain():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V13 Evidence", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(
            db=db,
            project_id=project.id,
            title="Missing security header: X-Content-Type-Options",
            severity="low",
            target="https://example.test/",
            description="Header absent.",
            recommendation="Add nosniff.",
            source="headers_check",
            evidence_kind="http_response_headers",
            evidence_content={"content-type": "text/html"},
        )
        evidence = db.query(Evidence).filter(Evidence.finding_id == finding.id).one()
        payload = evidence_chain_payload(db, evidence)
        assert len(payload["content_sha256"]) == 64
        assert len(payload["integrity_sha256"]) == 64
        assert evidence.project_id == project.id
        assert any(x["entity_type"] == "Finding" for x in payload["links"])
        assert db.query(EvidenceProvenance).filter(EvidenceProvenance.evidence_id == evidence.id).count() >= 1

        old_integrity = evidence.integrity_sha256
        evidence.content = '{"content-type":"text/html","x":"changed"}'
        db.commit(); db.refresh(evidence)
        assert evidence.integrity_sha256 != old_integrity
    finally:
        db.close()


def test_v13_finding_three_state_and_reopen_counter():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V13 Finding States", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(
            db=db,
            project_id=project.id,
            title="Authorization review candidate",
            severity="medium",
            target="https://example.test/api/profile/1001",
            description="Needs review.",
            recommendation="Enforce server-side authorization.",
            source="authorization_testing",
            evidence_kind="authorization_candidate_summary",
            evidence_content={"candidate": True},
        )
        assert finding.finding_state == "candidate"

        finding, lifecycle = update_finding_states(
            db,
            finding,
            finding_state="confirmed",
            remediation_state="remediation",
            verification_state="reproduced",
        )
        assert finding.finding_state == "confirmed"
        assert finding.verification_state == "reproduced"
        assert lifecycle.status == "remediation"

        finding, lifecycle = update_finding_states(
            db,
            finding,
            verification_state="resolved",
            remediation_state="resolved",
        )
        assert finding.verification_state == "resolved"
        assert lifecycle.status == "resolved"

        finding, lifecycle = update_finding_states(
            db,
            finding,
            verification_state="reproduced",
        )
        assert finding.reopened_count == 1
        assert lifecycle.status == "remediation"
    finally:
        db.close()


def test_v13_passive_nmap_burp_nuclei_imports_are_scope_filtered():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V13 Passive Import", scope_text="127.0.0.1")
        db.add(project); db.commit(); db.refresh(project)

        nmap = (
            b'<?xml version="1.0"?><nmaprun>'
            b'<host><status state="up"/><address addr="127.0.0.1" addrtype="ipv4"/>'
            b'<ports><port protocol="tcp" portid="443"><state state="open"/>'
            b'<service name="https" product="nginx" version="1.x"/></port></ports></host>'
            b'<host><status state="up"/><address addr="10.99.99.99" addrtype="ipv4"/></host>'
            b'</nmaprun>'
        )
        batch_nmap = import_nmap_xml(db, project, "scan.xml", nmap)
        assert batch_nmap.status == "done"
        assert batch_nmap.records_imported == 1
        assert batch_nmap.records_skipped == 1
        asset = db.query(Asset).filter(Asset.project_id == project.id, Asset.target == "127.0.0.1").one()
        assert db.query(Service).filter(Service.asset_id == asset.id, Service.port == 443).count() == 1

        raw_request = "GET /profile?lang=zh HTTP/1.1\r\nHost: 127.0.0.1\r\nAuthorization: Bearer BURP-SECRET\r\n\r\n"
        encoded = base64.b64encode(raw_request.encode()).decode()
        burp = (
            "<items><item>"
            "<url>http://127.0.0.1/profile?lang=zh</url>"
            "<host>127.0.0.1</host><port>80</port><protocol>http</protocol><path>/profile?lang=zh</path>"
            f'<request base64="true">{encoded}</request>'
            "</item></items>"
        ).encode()
        batch_burp = import_burp_xml(db, project, "burp.xml", burp)
        assert batch_burp.status == "done"
        stored = db.query(StoredRequest).filter(StoredRequest.project_id == project.id, StoredRequest.source == "burp_import").one()
        assert stored.method == "GET"
        assert stored.secret_headers_encrypted

        nuclei_line = json.dumps({
            "template-id": "missing-hsts",
            "matched-at": "http://127.0.0.1/",
            "info": {
                "name": "Missing HSTS",
                "severity": "low",
                "description": "Strict-Transport-Security not observed.",
                "classification": {"cwe-id": ["CWE-693"]},
            },
            "type": "http",
        })
        out_scope_line = json.dumps({
            "template-id": "outside",
            "matched-at": "http://10.99.99.99/",
            "info": {"name": "Outside", "severity": "high"},
        })
        batch_nuclei = import_nuclei_jsonl(
            db, project, "nuclei.jsonl",
            (nuclei_line + "\n" + out_scope_line).encode(),
        )
        assert batch_nuclei.records_imported == 1
        assert batch_nuclei.records_skipped == 1
        assert db.query(ImportBatch).filter(ImportBatch.project_id == project.id).count() == 3
        assert db.query(ImportRecord).count() >= 4
    finally:
        db.close()


def test_v13_txb02_word_report_is_valid_and_redacts_common_secrets():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(
            name="V13 Word Report",
            scope_text="example.test",
            client_name="测试客户",
            environment="test",
            status="testing",
        )
        db.add(project); db.commit(); db.refresh(project)
        create_finding(
            db=db,
            project_id=project.id,
            title="敏感接口安全验证",
            severity="medium",
            target="https://example.test/api/profile?ticket=URL-SECRET",
            description="用于 Word 报告 QA。",
            recommendation="修复后复测。",
            source="authorization_testing",
            evidence_kind="poc_request",
            evidence_content=(
                "GET /api/profile?ticket=URL-SECRET HTTP/1.1\n"
                "Host: example.test\n"
                "Authorization: Bearer HEADER-SECRET\n"
                "Cookie: sid=COOKIE-SECRET\n"
                '{"password":"BODY-SECRET"}'
            ),
            txb02_category="权限控制",
        )
        payload = generate_txb02_docx(db, project)
        assert payload[:2] == b"PK"
        doc = Document(io.BytesIO(payload))
        report_text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        )
        assert "渗透测试报告" in report_text
        assert "漏洞标题" in report_text
        assert "漏洞POC请求" in report_text
        assert "Evidence #" in report_text
        for secret in ["URL-SECRET", "HEADER-SECRET", "COOKIE-SECRET", "BODY-SECRET"]:
            assert secret not in report_text
    finally:
        db.close()


def test_v13_v12_database_migrates_to_v13_head(tmp_path):
    db_path = tmp_path / "v12.db"
    bootstrap = (
        "import sqlite3\n"
        f"db=sqlite3.connect(r'{db_path}')\n"
        "db.executescript(\"\"\""
        "CREATE TABLE persistent_jobs ("
        "id INTEGER PRIMARY KEY, project_id INTEGER, kind VARCHAR(100), target VARCHAR(1200), "
        "payload_json TEXT, scope_snapshot_json TEXT, scope_hash VARCHAR(80), status VARCHAR(60), "
        "priority INTEGER, attempts INTEGER, max_attempts INTEGER, timeout_seconds INTEGER, "
        "error TEXT, result_json TEXT, next_run_at DATETIME, started_at DATETIME, finished_at DATETIME, "
        "updated_at DATETIME, created_at DATETIME);"
        "CREATE TABLE evidence ("
        "id INTEGER PRIMARY KEY, finding_id INTEGER, task_id INTEGER, kind VARCHAR(100), content TEXT, created_at DATETIME);"
        "CREATE TABLE findings ("
        "id INTEGER PRIMARY KEY, project_id INTEGER, title VARCHAR(300), severity VARCHAR(30), target VARCHAR(1000), "
        "description TEXT, recommendation TEXT, source VARCHAR(100), fingerprint VARCHAR(80), dedupe_status VARCHAR(40), "
        "duplicate_of_id INTEGER, vuln_type VARCHAR(160), parameter VARCHAR(300), cwe_id VARCHAR(32), "
        "owasp_category VARCHAR(80), txb02_category VARCHAR(180), verification_state VARCHAR(60), created_at DATETIME);"
        "CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL);"
        "INSERT INTO alembic_version(version_num) VALUES('v1_2_core');"
        "\"\"\")\n"
        "db.commit(); db.close()\n"
    )
    subprocess.run([sys.executable, "-c", bootstrap], check=True)

    env = os.environ.copy()
    env["DATABASE_URL"] = f"sqlite:///{db_path}"
    code = (
        "from sqlalchemy import inspect\n"
        "from app.schema_migrations import ensure_schema_current\n"
        "from app.db import engine\n"
        "result=ensure_schema_current()\n"
        "i=inspect(engine)\n"
        "tables=set(i.get_table_names())\n"
        "job_cols={c['name'] for c in i.get_columns('persistent_jobs')}\n"
        "evidence_cols={c['name'] for c in i.get_columns('evidence')}\n"
        "finding_cols={c['name'] for c in i.get_columns('findings')}\n"
        "assert 'worker_id' in job_cols\n"
        "assert 'integrity_sha256' in evidence_cols\n"
        "assert 'finding_state' in finding_cols\n"
        "assert {'job_workers','evidence_provenance','import_batches','import_records'} <= tables\n"
        "with engine.connect() as conn: rev=conn.exec_driver_sql('SELECT version_num FROM alembic_version').scalar()\n"
        "assert rev == 'v1_8_toolchain_runs'\n"
        "print(result, rev)\n"
    )
    result = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(Path(__file__).resolve().parents[1]),
        env=env,
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "v1_8_toolchain_runs" in result.stdout

def test_v13_routes_worker_import_evidence_and_word_report():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V13 UI", scope_text="127.0.0.1")
        db.add(project); db.commit(); db.refresh(project)
        create_finding(
            db=db,
            project_id=project.id,
            title="Missing security header: X-Frame-Options",
            severity="low",
            target="http://127.0.0.1/",
            description="Header absent.",
            recommendation="Add a suitable framing policy.",
            source="headers_check",
            evidence_kind="http_response_headers",
            evidence_content={"content-type": "text/html"},
        )
        pid = project.id
    finally:
        db.close()

    with TestClient(app) as client:
        for url, marker in [
            (f"/projects/{pid}/jobs", "Worker 心跳"),
            (f"/projects/{pid}/imports", "被动结果导入"),
            (f"/projects/{pid}/evidence", "Content SHA256"),
            (f"/projects/{pid}/findings", "漏洞发现"),
            (f"/projects/{pid}/reports?legacy=1", "txb02 Word"),
        ]:
            response = client.get(url)
            assert response.status_code == 200
            assert marker in response.text
        docx = client.get(f"/projects/{pid}/reports/export.docx")
        assert docx.status_code == 200
        assert docx.content[:2] == b"PK"
        assert "wordprocessingml.document" in docx.headers["content-type"]


def test_v13_finding_evidence_redacts_common_secrets_before_persistence():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V13 Evidence Safety", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(
            db=db,
            project_id=project.id,
            title="POC Evidence Safety",
            severity="medium",
            target="https://example.test/profile?ticket=TARGET-SECRET",
            description="QA",
            recommendation="QA",
            source="authorization_testing",
            evidence_kind="poc_request",
            evidence_content=(
                "GET /profile?ticket=QUERY-SECRET HTTP/1.1\\n"
                "Host: example.test\\n"
                "Authorization: Bearer HEADER-SECRET\\n"
                "Cookie: sid=COOKIE-SECRET\\n"
                '{"password":"BODY-SECRET","safe":"visible"}'
            ),
        )
        evidence = db.query(Evidence).filter(Evidence.finding_id == finding.id).one()
        assert evidence.redaction_state == "redacted"
        assert "TARGET-SECRET" not in finding.target
        assert "••••" in finding.target
        for secret in ["QUERY-SECRET", "HEADER-SECRET", "COOKIE-SECRET", "BODY-SECRET"]:
            assert secret not in evidence.content
        assert '"safe":"visible"' in evidence.content or '"safe": "visible"' in evidence.content
        assert "••••" in evidence.content
    finally:
        db.close()
