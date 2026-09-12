import base64
import io
import json
import os
import subprocess
import sys
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.ai.factory import current_ai_provider_name
from app.db import Base, SessionLocal, engine
from app.main import app
from app.models import (
    AppPreference,
    AuthorizationMatrixRun,
    BackupRecord,
    EndpointParameter,
    Evidence,
    Project,
    ProjectSkill,
    SkillDefinition,
    StoredRequest,
)
from app.services.authorization_matrix import compare_matrix_runs
from app.services.evidence_attachments import save_finding_screenshot
from app.services.finding_service import create_finding
from app.services.passive_imports import import_har, import_openapi, import_postman_collection
from app.services.personal_backup import create_project_backup
from app.services.personal_diagnostics import diagnostic_snapshot
from app.services.personal_next_steps import suggested_next_steps
from app.services.personal_settings import get_personal_settings, save_personal_settings
from app.services.project_templates import apply_project_template, normalize_quick_target
from app.services.txb02_report import generate_txb02_docx

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


def test_v14_personal_settings_encrypt_ai_key_and_runtime_override():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        save_personal_settings(
            db,
            {"setup_completed": True, "default_template": "api-authz", "ai_provider": "openai_compatible", "ai_api_base": "https://ai.example.test/v1", "ai_model": "qa-model", "backup_enabled": False},
            ai_api_key="V14-API-SECRET",
        )
        prefs = get_personal_settings(db)
        assert prefs["setup_completed"] is True
        assert prefs["default_template"] == "api-authz"
        assert prefs["ai_api_key_configured"] is True
        row = db.query(AppPreference).filter(AppPreference.key == "personal_config").one()
        assert "V14-API-SECRET" not in row.value_json
        assert "V14-API-SECRET" not in row.secret_encrypted
    finally:
        db.close()
    assert current_ai_provider_name() == "openai_compatible"
    db = SessionLocal()
    try:
        save_personal_settings(db, {"ai_provider": "mock", "ai_api_base": "", "ai_model": "", "backup_enabled": True}, ai_api_key="")
    finally:
        db.close()
    assert current_ai_provider_name() == "mock"


def test_v14_project_templates_and_quick_target():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V14 Template", scope_text="api.example.test", template_slug="api-authz")
        db.add(project); db.commit(); db.refresh(project)
        result = apply_project_template(db, project.id, "api-authz")
        enabled = db.query(SkillDefinition.slug).join(ProjectSkill, ProjectSkill.skill_id == SkillDefinition.id).filter(ProjectSkill.project_id == project.id, ProjectSkill.enabled == 1).all()
        slugs = {x[0] for x in enabled}
        assert result["slug"] == "api-authz"
        assert {"authorization-differential-review", "request-response-diff", "agent-tool-guardrails"} <= slugs
    finally:
        db.close()
    target, scope = normalize_quick_target("api.example.test/path")
    assert target == "https://api.example.test/path"
    assert scope == "api.example.test"


def test_v14_har_postman_openapi_imports_create_endpoints_without_execution():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V14 Imports", scope_text="api.example.test")
        db.add(project); db.commit(); db.refresh(project)
        har = json.dumps({"log": {"entries": [
            {"request": {"method": "GET", "url": "https://api.example.test/users/1001?lang=zh", "headers": [{"name": "Accept", "value": "application/json"}]}},
            {"request": {"method": "GET", "url": "https://outside.example.test/", "headers": []}},
        ]}}).encode()
        bh = import_har(db, project, "traffic.har", har)
        assert (bh.records_imported, bh.records_skipped) == (1, 1)
        postman = json.dumps({"info": {"name": "QA"}, "variable": [{"key": "baseUrl", "value": "https://api.example.test"}], "item": [
            {"name": "Get profile", "request": {"method": "GET", "url": "{{baseUrl}}/profile?view=full", "header": [{"key": "Authorization", "value": "Bearer POSTMAN-SECRET"}]}},
            {"name": "Update profile", "request": {"method": "POST", "url": "{{baseUrl}}/profile", "header": [{"key": "Content-Type", "value": "application/json"}], "body": {"mode": "raw", "raw": "{\"nickname\":\"qa\"}"}}},
        ]}).encode()
        bp = import_postman_collection(db, project, "collection.json", postman)
        assert bp.records_imported == 2
        openapi = json.dumps({"openapi": "3.0.3", "servers": [{"url": "https://api.example.test"}], "paths": {
            "/orders/{orderId}": {"get": {"summary": "Order detail", "parameters": [{"name": "orderId", "in": "path", "required": True, "schema": {"type": "integer"}}, {"name": "verbose", "in": "query", "required": False, "schema": {"type": "boolean"}}]}},
            "/orders": {"post": {"summary": "Create order", "requestBody": {"content": {"application/json": {"schema": {"type": "object", "properties": {"productId": {"type": "integer"}, "token": {"type": "string"}}}}}}}},
        }}).encode()
        bo = import_openapi(db, project, "openapi.json", openapi)
        assert bo.records_imported == 2
        requests = db.query(StoredRequest).filter(StoredRequest.project_id == project.id).all()
        assert any(r.source == "har_import" for r in requests)
        assert any(r.source == "postman_import" and r.policy_class == "STATE_CHANGE" for r in requests)
        assert any(r.source == "openapi_import" and r.method == "POST" and r.policy_class == "STATE_CHANGE" for r in requests)
        params = db.query(EndpointParameter).all()
        order_id = next(p for p in params if p.name == "orderId")
        assert order_id.location == "path" and order_id.required == 1
        token = next(p for p in params if p.name == "token")
        assert token.sensitive == 1 and token.example_redacted == "••••"
    finally:
        db.close()


def test_v14_screenshot_evidence_enters_word_report(tmp_path, monkeypatch):
    Base.metadata.create_all(bind=engine)
    monkeypatch.setattr("app.services.evidence_attachments.settings.evidence_artifact_dir", str(tmp_path / "evidence"))
    db = SessionLocal()
    try:
        project = Project(name="V14 Screenshot", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        finding = create_finding(db=db, project_id=project.id, title="Screenshot QA Finding", severity="medium", target="https://example.test/profile?ticket=SCREEN-SECRET", description="QA", recommendation="QA", source="authorization_testing", evidence_kind="authorization_candidate_summary", evidence_content={"candidate": True})
        attachment = save_finding_screenshot(db, finding, "切换普通用户后仍可访问", "image/png", PNG_1X1)
        assert len(attachment.file_sha256) == 64 and Path(attachment.file_path).exists()
        ev = db.get(Evidence, attachment.evidence_id)
        assert ev.kind == "screenshot_evidence" and len(ev.content_sha256) == 64
        payload = generate_txb02_docx(db, project)
        assert payload[:2] == b"PK"
        doc = Document(io.BytesIO(payload))
        text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
        assert "验证截图" in text and "切换普通用户后仍可访问" in text and "SCREEN-SECRET" not in text
        import zipfile
        with zipfile.ZipFile(io.BytesIO(payload)) as zf:
            assert any(name.startswith("word/media/") for name in zf.namelist())
    finally:
        db.close()


def test_v14_backup_retention_and_diagnostics(tmp_path):
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        save_personal_settings(db, {"backup_enabled": True, "backup_dir": str(tmp_path / "backups"), "backup_retention": 2, "backup_interval_hours": 24, "ai_provider": "mock"})
        project = Project(name="V14 Backup", scope_text="example.test")
        db.add(project); db.commit(); db.refresh(project)
        for _ in range(3): create_project_backup(db, project, backup_type="sanitized_manual")
        rows = db.query(BackupRecord).filter(BackupRecord.project_id == project.id).order_by(BackupRecord.id.desc()).all()
        assert len(rows) == 2 and all(Path(x.file_path).exists() for x in rows)
        snap = diagnostic_snapshot(db)
        assert snap["total"] >= 8
        assert any(x["name"] == "数据库" and x["ok"] for x in snap["checks"])
        assert any(x["name"] == "备份目录" and x["ok"] for x in snap["checks"])
    finally:
        db.close()


def test_v14_personal_next_steps_prioritize_identity_and_authorization():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V14 Next Steps", scope_text="api.example.test")
        db.add(project); db.commit(); db.refresh(project)
        db.add(StoredRequest(project_id=project.id, name="GET profile", method="GET", url="https://api.example.test/profile", headers_json="{}", secret_headers_encrypted="", body="", source="manual", policy_class="READ_ONLY")); db.commit()
        steps = suggested_next_steps(db, project)
        assert any("测试身份" in x["title"] for x in steps)
    finally:
        db.close()


def test_v14_authorization_matrix_history_comparison():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        project = Project(name="V14 Matrix Compare", scope_text="api.example.test")
        db.add(project); db.commit(); db.refresh(project)
        left = AuthorizationMatrixRun(project_id=project.id, stored_request_id=100, status="done", summary_json=json.dumps({"stored_request_id": 100, "comparisons": [{"comparison_identity_id": 2, "comparison_name": "User B", "classification": "potential_horizontal_authorization_gap", "confidence": 95}]}))
        right = AuthorizationMatrixRun(project_id=project.id, stored_request_id=100, status="done", summary_json=json.dumps({"stored_request_id": 100, "comparisons": [{"comparison_identity_id": 2, "comparison_name": "User B", "classification": "authorization_control_enforced", "confidence": 98}]}))
        db.add_all([left,right]); db.commit(); db.refresh(left); db.refresh(right)
        diff = compare_matrix_runs(left,right)
        assert diff["compatible"] is True and diff["improved"] == 1 and diff["regressed"] == 0 and diff["rows"][0]["delta"] == "improved"
    finally:
        db.close()


def test_v14_v13_database_migrates_to_v14_head(tmp_path):
    db_path = tmp_path / "v13.db"
    sql = ("CREATE TABLE projects (id INTEGER PRIMARY KEY, name VARCHAR(200), scope_text TEXT, client_name VARCHAR(240), environment VARCHAR(80), engagement_type VARCHAR(80), status VARCHAR(60), start_date VARCHAR(32), end_date VARCHAR(32), authorization_note TEXT, testers_json TEXT, created_at DATETIME);" "CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL);" "INSERT INTO alembic_version(version_num) VALUES('v1_3_enterprise');")
    bootstrap = "import sqlite3\n" + f"db=sqlite3.connect(r'{db_path}')\n" + f"db.executescript({sql!r})\n" + "db.commit();db.close()\n"
    subprocess.run([sys.executable, "-c", bootstrap], check=True)
    env = os.environ.copy(); env["DATABASE_URL"] = f"sqlite:///{db_path}"
    code = ("from sqlalchemy import inspect\nfrom app.db import engine\nfrom app.schema_migrations import ensure_schema_current\nensure_schema_current()\ni=inspect(engine)\ntables=set(i.get_table_names())\ncols={c['name'] for c in i.get_columns('projects')}\nassert 'template_slug' in cols\nassert {'app_preferences','evidence_attachments','backup_records'} <= tables\nwith engine.connect() as c: rev=c.exec_driver_sql('SELECT version_num FROM alembic_version').scalar()\nassert rev=='v1_6_2_assetux'\nprint(rev)\n")
    result = subprocess.run([sys.executable, "-c", code], cwd=str(Path(__file__).resolve().parents[1]), env=env, capture_output=True, text=True, timeout=120)
    assert result.returncode == 0, result.stdout + result.stderr
    assert "v1_6_2_assetux" in result.stdout


def test_v14_quick_scan_setup_and_core_pages(monkeypatch):
    Base.metadata.create_all(bind=engine)
    async def no_network_job(job_id): return None
    monkeypatch.setattr("app.personal_routes.run_job_now", no_network_job)
    with TestClient(app) as client:
        setup = client.get("/setup"); assert setup.status_code == 200 and "首次启动向导" in setup.text
        diagnostics = client.get("/diagnostics"); assert diagnostics.status_code == 200 and "系统诊断" in diagnostics.text
        home = client.get("/"); assert home.status_code == 200 and "QUICK SCAN" in home.text and "单目标快速测试" in home.text
        response = client.post("/quick-scan", data={"target":"https://quick.example.test/app","template_slug":"quick-web","project_name":"Quick QA"}, follow_redirects=False)
        assert response.status_code == 303 and response.headers["location"].startswith("/projects/")
    db=SessionLocal()
    try:
        project=db.query(Project).filter(Project.name=="Quick QA").order_by(Project.id.desc()).first()
        assert project is not None and project.scope_text=="quick.example.test" and project.template_slug=="quick-web"
    finally: db.close()


def test_v14_personal_launcher_exists():
    root=Path(__file__).resolve().parents[1]
    launcher=(root/"START_PERSONAL.bat").read_text(encoding="utf-8")
    assert "worker.py" in launcher.lower()
    assert "scripts\\db-upgrade.py" in launcher
    assert "http://127.0.0.1:8000/" in launcher
