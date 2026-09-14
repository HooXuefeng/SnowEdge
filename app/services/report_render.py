"""Real document exports of the existing immutable delivery projection."""
import base64,csv,html,io,json,math,hashlib,threading
from collections import OrderedDict
from pathlib import Path
from ..ui_i18n import SEVERITY_ZH,STATUS_ZH
from .report_data import TEMPLATES,STAGES

LABELS={'target':'漏洞 URL','parameter':'漏洞参数','vuln_type':'漏洞类型','cwe_id':'CWE','owasp_category':'OWASP 分类','difficulty':'利用难度','cvss':'CVSS','created_at':'发现时间','affected_assets':'影响资产'}
ROLES={'normal':'正常验证','poc':'POC 验证','baseline':'基准身份','comparison':'对照身份','retest':'复测'}
def text(value):
    if isinstance(value,list) and all(isinstance(x,str) for x in value):return '、'.join(value) or '未记录'
    return json.dumps(value,ensure_ascii=False,indent=2) if isinstance(value,(dict,list)) else str(value or '未记录')

def timestamp(value):
    return str(value or '未记录').split('.')[0].replace('T',' ')

def blocks(data):
    info=data['project_info'];summary=data['summary'];c=data['coverage'];findings=data['findings']
    brief=data.get('template')=='brief'
    retest=data.get('template')=='retest' or data.get('stage')=='retest'
    yield {'kind':'heading','level':1,'text':'1 项目概况','anchor':'overview'}
    yield {'kind':'table','caption':'表 1 项目概况','headers':['项目','内容'],'rows':[[label,text(value)] for label,value in [('测试目标',info['objectives']),('客户 / 目标',info['client_name']),('测试范围',info['scope_text']),('测试方式',info['method']),('测试时间',info['start_date']+' 至 '+info['end_date']),('测试人员',info['testers']),('测试限制',info['limitations']),('授权说明',info['authorization_note'])]]}
    yield {'kind':'heading','level':1,'text':'2 测试结果摘要','anchor':'summary'}
    yield {'kind':'paragraph','text':summary['conclusion']}
    yield {'kind':'chart','counts':summary['severity']}
    yield {'kind':'paragraph','text':f"已修复 {summary['fixed']} 项；未修复 {summary['unfixed']} 项；待复测 {summary['awaiting_retest']} 项。统计对应本报告选中的漏洞。"}
    if data['warnings']:
        yield {'kind':'heading','level':2,'text':'交付核对事项'}
        for warning in dict.fromkeys(data['warnings']):yield {'kind':'paragraph','text':warning}
    yield {'kind':'heading','level':1,'text':'3 测试覆盖情况','anchor':'coverage'}
    yield {'kind':'table','caption':'表 2 测试覆盖','headers':['指标','数量'],'rows':[['已登记 / 已验证资产',f"{c['assets']} / {c['tested_assets']}"],['已登记 / 已验证接口',f"{c['endpoints']} / {c['tested_endpoints']}"],['身份数量',str(c['identities'])],['权限场景 / 已完成',f"{c['permission_scenarios']} / {c['permission_completed']}"],['接口回放覆盖率',str(c['rate'])+'%' if c['rate'] is not None else '无分母，暂不可计算']]}
    yield {'kind':'paragraph','text':c['basis']}
    yield {'kind':'paragraph','text':'已测试功能模块：'+text(info['modules'])}
    yield {'kind':'paragraph','text':'未覆盖接口：\n'+('\n'.join(c['uncovered']) or '当前已登记接口无未回放项；未登记区域无法据此判断。')}
    yield {'kind':'heading','level':1,'text':'4 漏洞汇总','anchor':'findings'}
    yield {'kind':'table','caption':'表 3 漏洞汇总','headers':['编号 / 名称','位置 / 资产','类型 / 难度','风险','状态 / 复测'],'rows':[[f['number']+' '+f['title'],f['target']+'\n'+f['affected_assets'],text(f['vuln_type'])+'\n'+text(f['difficulty']),SEVERITY_ZH.get(f['severity'],f['severity']),f['status_label']+'\n'+STATUS_ZH.get(f['verification_state'],f['verification_state'])] for f in findings]}
    if data.get('template')=='scan':
        yield {'kind':'paragraph','text':'扫描报告中的条目以保存的发现和证据为准。自动探测结果不代表已完成人工利用验证；确认状态与验证过程应分别核对。'}
    yield {'kind':'heading','level':1,'text':'5 漏洞详情','anchor':'details','page_break':bool(findings)}
    for index,f in enumerate(findings,1):
        yield {'kind':'heading','level':1,'text':f"5.{index} {f['number']} {f['title']}",'anchor':f"finding-{f['id']}",'page_break':index>1}
        yield {'kind':'table','caption':f'表 4-{index} 漏洞属性','headers':['字段','内容'],'rows':[[label,(timestamp(f.get(key)) if key=='created_at' else text(f.get(key)))] for key,label in LABELS.items()]+[['风险等级',SEVERITY_ZH.get(f['severity'],f['severity'])],['漏洞状态',f['status_label']],['复测状态',STATUS_ZH.get(f['verification_state'],f['verification_state'])]]}
        for key,label in [('description','漏洞描述'),('principle','漏洞原理'),('impact','影响分析'),('verification_steps','验证过程')]:
            yield {'kind':'heading','level':2,'text':label};yield {'kind':'paragraph','text':text(f.get(key))}
        if retest:
            yield {'kind':'heading','level':2,'text':'整改与复测结论'}
            yield {'kind':'paragraph','text':f['status_label']+'；'+STATUS_ZH.get(f['verification_state'],f['verification_state'])+'。初测与复测证据均保留在本章节，历史交付版本保持不变。'}
        if brief:
            yield {'kind':'paragraph','text':'简版省略原始 HTTP 与证据正文；完整数据保留在同版本 JSON 文件中。'}
        for h in ([] if brief else f['http']):
            label=ROLES.get(h['role'],h['role'])
            yield {'kind':'heading','level':2,'text':label+' / '+h['identity']}
            yield {'kind':'paragraph','text':h['source_note']+f" 请求 #{h['request_id']} / 回放 #{h['replay_id'] or '未保存'}"}
            for part in ('request','response'):yield {'kind':'http','caption':label+('请求' if part=='request' else '响应'),'text':h[part] or '尚未保存响应，不能补造。'}
        for comparison in ([] if brief else f.get('permission_comparisons',[])):
            yield {'kind':'heading','level':2,'text':f"权限差异验证 #{comparison['case_id']}"}
            for role in ('baseline','comparison'):
                if role in comparison:
                    h=comparison[role];yield {'kind':'paragraph','text':ROLES[role]+'：'+h['identity']}
                    yield {'kind':'http','caption':'请求','text':h['request']};yield {'kind':'http','caption':'响应','text':h['response'] or '未保存'}
            yield {'kind':'paragraph','text':'权限差异：'+text(comparison.get('difference',comparison['summary']))+'\n验证结果：'+comparison['conclusion']}
        yield {'kind':'heading','level':2,'text':'验证证据'}
        figure_number=0;shown_images=set()
        for e in f['evidence']:
            yield {'kind':'paragraph','text':f"{e['caption']} · 证据 #{e['id']} · {timestamp(e['created_at'])}\n{e['description']}"}
            content=e['content']
            if not brief:
                if isinstance(content,dict) and content.get('report_evidence'):
                    for key,label in [('request','请求'),('response','响应')]:
                        if content.get(key):yield {'kind':'http','caption':label,'text':content[key]}
                    if content.get('note'):yield {'kind':'paragraph','text':text(content['note'])}
                elif e['kind'] not in ('screenshot','finding_screenshot','screenshot_evidence'):
                    yield {'kind':'http' if e['kind'] in ('http_request','http_response','console','code') else 'paragraph','caption':e['caption'],'text':text(content)}
            for n,image in enumerate(f['images']):
                if image.get('evidence_id') in (e['id'],e.get('source_id')):
                    figure_number+=1;shown_images.add(n)
                    yield {'kind':'image','data':image['data'],'caption':f"图 {index}-{figure_number} {e['caption']}"}
        for n,image in enumerate(f['images']):
            if n not in shown_images:
                figure_number+=1
                yield {'kind':'image','data':image['data'],'caption':f"图 {index}-{figure_number} {image['label']}"}
        for key,label in [('recommendation','修复建议'),('references','参考资料')]:
            yield {'kind':'heading','level':2,'text':label};yield {'kind':'paragraph','text':text(f.get(key))}
        yield {'kind':'heading','level':2,'text':'复测结果与时间线'}
        for e in f['retests']:
            content=e['content'] if isinstance(e['content'],dict) else {}
            conclusion=e.get('conclusion') or content.get('conclusion','')
            yield {'kind':'paragraph','text':timestamp(e['created_at'])+' · '+e['caption']+'\n复测结论：'+STATUS_ZH.get(conclusion,conclusion or '未记录')+f"。请求、响应及截图见本章证据 #{e['id']}。"}
        if not f['retests']:yield {'kind':'paragraph','text':'尚未关联复测证据。'}

def http_chunks(value,max_rows=28):
    chunk=[];rows=0
    for line in value.splitlines(keepends=True):
        cost=max(1,math.ceil(len(line)/80))
        if rows+cost>max_rows and chunk:yield ''.join(chunk);chunk=[];rows=0
        chunk.append(line);rows+=cost
    if chunk:yield ''.join(chunk)

def html_report(data,toc_pages=None):
    esc=html.escape;parts=[];toc=[]
    for b in blocks(data):
        kind=b['kind']
        if kind=='heading':
            if b['level']==1:toc.append((b['anchor'],b['text']))
            parts.append(f"<h{b['level']} id='{b.get('anchor','')}' class='{'new-page' if b.get('page_break') else ''}'>{esc(b['text'])}</h{b['level']}>")
        elif kind=='table':parts.append('<div class="caption">'+esc(b['caption'])+'</div><table><thead><tr>'+''.join('<th>'+esc(x)+'</th>' for x in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+esc(x)+'</td>' for x in row)+'</tr>' for row in b['rows'])+'</tbody></table>')
        elif kind=='http':
            for n,chunk in enumerate(http_chunks(b['text'])):parts.append('<div class="http"><div class="caption">'+esc(b.get('caption','HTTP'))+('（续）' if n else '')+'</div><pre>'+esc(chunk)+'</pre></div>')
        elif kind=='image':parts.append('<figure><img src="data:image/png;base64,'+b['data']+'"><figcaption>'+esc(b['caption'])+'</figcaption></figure>')
        elif kind=='chart':
            parts.append('<div class="chart">');total=max(1,sum(b['counts'].values()))
            colors={'critical':'#991b1b','high':'#dc2626','medium':'#d97706','low':'#2563eb','info':'#64748b'}
            for severity in colors:
                count=b['counts'].get(severity,0);parts.append(f"<div>{esc(SEVERITY_ZH.get(severity,severity))} {count}<span style='width:{100*count/total}%;background:{colors[severity]}'></span></div>")
            parts.append('</div>')
        else:parts.append('<p>'+esc(b['text'])+'</p>')
    logo=Path(__file__).resolve().parents[1]/'static/brand/snowedge-app.png'
    mark='<img class="brand" src="data:image/png;base64,'+base64.b64encode(logo.read_bytes()).decode()+'">' if logo.exists() else ''
    css=Path(__file__).resolve().parents[1].joinpath('static/report-document.css').read_text(encoding='utf-8')
    cover=f"<section class='cover'>{mark}<p>SnowEdge</p><h1>{esc(data['project'])}</h1><h2>{esc(TEMPLATES[data['template']])}</h2><p>{esc(STAGES[data['stage']])}</p><p>客户 / 目标：{esc(data['project_info']['client_name'] or '未记录')}</p><p>测试时间：{esc(data['project_info']['start_date'])} 至 {esc(data['project_info']['end_date'])}</p><p>报告版本：{esc(data['report_version'])}<br>生成时间：{esc(timestamp(data['created_at']))} UTC</p></section>"
    toc_html='<section class="toc"><h1>目录</h1>'+''.join(f'<a href="#{a}">{esc(t)}<span class="toc-page">{(toc_pages or {}).get(a,"")}</span></a>' for a,t in toc)+'</section>'
    return '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; img-src data:; style-src \'unsafe-inline\'"><title>'+esc(data['project'])+'</title><style>'+css+'</style></head><body><main>'+cover+toc_html+''.join(parts)+'</main></body></html>'

_pdf_cache=OrderedDict()
_pdf_lock=threading.RLock()
_pdfium_lock=threading.Lock()

def pdf_report(data):
    key=hashlib.sha256(json.dumps(data,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
    with _pdf_lock:
        if key in _pdf_cache:
            _pdf_cache.move_to_end(key);return _pdf_cache[key]
        result=_generate_pdf(data)
        if len(result)<=32*1024*1024:
            while _pdf_cache and (len(_pdf_cache)>=3 or sum(map(len,_pdf_cache.values()))+len(result)>32*1024*1024):_pdf_cache.popitem(last=False)
            _pdf_cache[key]=result
        return result

def preview_pages(data,page_number=None):
    try:import pypdfium2 as pdfium
    except ImportError:raise ValueError('缺少 PDF 预览组件，请更新 requirements.txt 中的依赖后重试。') from None
    raw=pdf_report(data)
    with _pdfium_lock:
        with pdfium.PdfDocument(raw) as document:
            if page_number is None:return {'pages':len(document)}
            if not 1<=page_number<=len(document):raise ValueError('页码超出报告范围')
            page=document[page_number-1]
            try:
                bitmap=page.render(scale=1.6)
                try:
                    image=bitmap.to_pil();output=io.BytesIO();image.save(output,format='PNG');return output.getvalue()
                finally:bitmap.close()
            finally:page.close()

def _generate_pdf(data):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser=None
        for options in ({'channel':'chrome'},{'channel':'msedge'},{}):
            try:browser=p.chromium.launch(headless=True,**options);break
            except Exception:continue
        if browser is None:raise ValueError('PDF 引擎不可用，请安装 Chrome/Edge，或执行 python -m playwright install chromium。未生成文件。')
        try:
            page=browser.new_page();page.route('**/*',lambda route:route.abort())
            page.set_content(html_report(data),wait_until='load');page.evaluate('document.fonts.ready')
            options=dict(format='A4',print_background=True,display_header_footer=True,header_template='<div style="width:100%;margin:0 16mm;font-size:9px;color:#64748b">SnowEdge · 安全测试报告</div>',footer_template='<div style="width:100%;margin:0 16mm;font-size:9px;color:#64748b;text-align:right">第 <span class="pageNumber"></span> 页 / 共 <span class="totalPages"></span> 页</div>',prefer_css_page_size=True)
            result=page.pdf(**options)
            try:import pypdfium2 as pdfium
            except ImportError:return result
            anchors=[(b['anchor'],''.join(b['text'].split())) for b in blocks(data) if b['kind']=='heading' and b['level']==1]
            locations={}
            with _pdfium_lock:
                with pdfium.PdfDocument(result) as document:
                    for index in range(2,len(document)):
                        pdf_page=document[index]
                        try:
                            textpage=pdf_page.get_textpage()
                            try:body=''.join(textpage.get_text_range().split())
                            finally:textpage.close()
                            for anchor,label in anchors:
                                if anchor not in locations and label in body:locations[anchor]=index+1
                        finally:pdf_page.close()
            page.set_content(html_report(data,locations),wait_until='load');page.evaluate('document.fonts.ready')
            return page.pdf(**options)
        finally:browser.close()

def docx_report(data):
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(20);s.left_margin=s.right_margin=Mm(16)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        style=d.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    d.styles['Normal'].font.size=Pt(10)
    d.styles['Normal'].paragraph_format.line_spacing=1.2
    d.styles['Caption'].paragraph_format.keep_with_next=True
    for name,size in [('Title',26),('Heading 1',16),('Heading 2',12),('Subtitle',14)]:d.styles[name].font.size=Pt(size)
    for name in ('Title','Subtitle'):
        style=d.styles[name]
        for border in style.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    d.styles['Subtitle'].font.italic=False
    s.header.paragraphs[0].text='SnowEdge · 安全测试报告';footer=s.footer.paragraphs[0];footer.alignment=2;footer.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field);footer.add_run(' 页 / 共 ');total=OxmlElement('w:fldSimple');total.set(qn('w:instr'),'NUMPAGES');footer._p.append(total);footer.add_run(' 页')
    logo=Path(__file__).resolve().parents[1]/'static/brand/snowedge-app.png'
    if logo.exists():d.add_picture(str(logo),width=Mm(24))
    d.add_paragraph(data['project'],style='Title');d.add_paragraph(TEMPLATES[data['template']],style='Subtitle')
    for label,value in [('客户',data['project_info']['client_name']),('阶段',STAGES[data['stage']]),('测试时间',data['project_info']['start_date']+' 至 '+data['project_info']['end_date']),('报告版本',data['report_version']),('生成时间',timestamp(data['created_at'])+' UTC')]:d.add_paragraph(label+'：'+text(value))
    d.add_page_break();d.add_paragraph('目录',style='Title')
    paragraph=d.add_paragraph();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');paragraph.add_run()._r.append(begin)
    instruction=OxmlElement('w:instrText');instruction.set(qn('xml:space'),'preserve');instruction.text=' TOC \\o "1-1" \\h \\z \\u ';paragraph.add_run()._r.append(instruction)
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate');paragraph.add_run()._r.append(separate)
    for b in blocks(data):
        if b['kind']=='heading' and b['level']==1:d.add_paragraph(b['text'])
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');d.add_paragraph().add_run()._r.append(end)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');d.settings.element.append(update)
    d.add_page_break()
    for b in blocks(data):
        kind=b['kind']
        if kind=='heading':
            if b.get('page_break'):d.add_page_break()
            d.add_heading(b['text'],b['level'])
        elif kind=='table':
            d.add_paragraph(b['caption'],style='Caption');table=d.add_table(rows=1,cols=len(b['headers']));table.style='Table Grid';table.autofit=False
            for i,x in enumerate(b['headers']):table.rows[0].cells[i].text=x
            repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in b['rows']:
                cells=table.add_row().cells
                for i,x in enumerate(row):cells[i].text=x
            for row in table.rows:
                prevent=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(prevent)
                for cell in row.cells:cell.width=Mm(178/len(b['headers']))
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'DCE3ED');borders.append(border)
            table._tbl.tblPr.append(borders)
            for cell in table.rows[0].cells:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEF2FF');cell._tc.get_or_add_tcPr().append(shade)
        elif kind=='image':
            from PIL import Image
            blob=base64.b64decode(b['data'])
            with Image.open(io.BytesIO(blob)) as im:width=min(178,185*im.width/im.height)
            d.add_picture(io.BytesIO(blob),width=Mm(width));d.paragraphs[-1].paragraph_format.keep_with_next=True;caption=d.add_paragraph(b['caption'],style='Caption');caption.paragraph_format.keep_with_next=False
        elif kind=='http':
            for i,chunk in enumerate(http_chunks(b['text'])):
                d.add_paragraph(b.get('caption','HTTP')+('（续）' if i else ''),style='Caption')
                paragraph=d.add_paragraph();paragraph.paragraph_format.keep_together=True
                paragraph.paragraph_format.line_spacing=1.1
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F5F7FA');paragraph._p.get_or_add_pPr().append(shade)
                run=paragraph.add_run(chunk.replace('\r\n','\n'));run.font.name='Consolas';run.font.size=Pt(8)
        elif kind=='chart':
            from PIL import Image,ImageDraw,ImageFont
            chart=Image.new('RGB',(1500,200),'white');draw=ImageDraw.Draw(chart)
            try:font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',30);labels=SEVERITY_ZH
            except OSError:font=ImageFont.load_default(size=30);labels={}
            total=max(1,sum(b['counts'].values()))
            for n,(key,color) in enumerate([('critical','#991b1b'),('high','#dc2626'),('medium','#d97706'),('low','#2563eb'),('info','#64748b')]):
                count=b['counts'].get(key,0);x=n*300
                draw.text((x+8,20),labels.get(key,key)+' '+str(count),font=font,fill='#172033')
                if count:draw.rectangle((x+8,95,x+8+int(270*count/total),125),fill=color)
            output=io.BytesIO();chart.save(output,format='PNG');d.add_picture(io.BytesIO(output.getvalue()),width=Mm(178))
        else:d.add_paragraph(b['text'])
    output=io.BytesIO();d.save(output);return output.getvalue()

def export(data,fmt):
    if fmt=='json':return json.dumps(data,ensure_ascii=False,indent=2).encode(),'application/json'
    if fmt=='csv':
        output=io.StringIO(newline='');w=csv.writer(output);w.writerow(['编号','漏洞名称','漏洞位置','漏洞类型','风险等级','利用难度','影响资产','漏洞状态','复测状态'])
        for f in data['findings']:
            values=[f['number'],f['title'],f['target'],f['vuln_type'],f['severity'],f['difficulty'],f['affected_assets'],f['status_label'],f['verification_state']]
            w.writerow(["'"+x if x.lstrip().startswith(('=','+','-','@')) else x for x in values])
        return output.getvalue().encode('utf-8-sig'),'text/csv; charset=utf-8'
    if fmt=='html':return html_report(data).encode(),'text/html; charset=utf-8'
    if fmt=='pdf':return pdf_report(data),'application/pdf'
    if fmt=='docx':return docx_report(data),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if fmt=='md':return '\n\n'.join(('#'*b['level']+' '+b['text']) if b['kind']=='heading' else b.get('text',text(b)) for b in blocks(data)).encode(),'text/markdown; charset=utf-8'
    raise ValueError('不支持的报告格式')
