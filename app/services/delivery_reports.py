"""Frozen delivery content: later edits cannot alter an existing report."""
import base64
import hashlib
import html
import io
import json
import secrets
from ..models import Finding, Evidence, EvidenceAttachment, AppPreference, _utcnow
from ..ui_i18n import STATUS_ZH, SEVERITY_ZH
from .evidence_safety import redact_object
from .evidence_attachments import render_annotated_image_bytes

VARIANTS={'client':'客户版','internal':'内部技术版','retest':'复测版'}

def encoded(value):return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(',',':'))

def create_snapshot(db,project,ids,variant,include_images=False):
    if variant not in VARIANTS:raise ValueError('报告类型无效。')
    if not ids or len(ids)>200:raise ValueError('请选择 1–200 条漏洞。')
    rows=db.query(Finding).filter(Finding.project_id==project.id,Finding.id.in_(ids)).order_by(Finding.id).all()
    if {r.id for r in rows}!=set(ids):raise ValueError('所选漏洞不存在或不属于当前项目。')
    data={'project':project.name,'variant':variant,'created_at':str(_utcnow()),'findings':[],'warnings':[]}
    image_total=0
    for row in rows:
        item={key:getattr(row,key) for key in ('id','title','severity','target','description','recommendation','vuln_type','parameter','cwe_id','txb02_category','owasp_category','finding_state','verification_state')}
        item,_=redact_object(item)
        evidence=db.query(Evidence).filter_by(project_id=project.id,finding_id=row.id).order_by(Evidence.id).all()
        item['evidence']=[];item['images']=[]
        if not evidence:data['warnings'].append(f'漏洞 #{row.id} 缺少关联证据。')
        if row.finding_state!='confirmed':data['warnings'].append(f'漏洞 #{row.id} 尚未人工确认为有效漏洞。')
        if not row.recommendation:data['warnings'].append(f'漏洞 #{row.id} 缺少修复建议。')
        if variant=='retest' and row.verification_state=='unverified':data['warnings'].append(f'漏洞 #{row.id} 尚无复测结论。')
        for e in evidence:
            entry={'id':e.id,'kind':e.kind,'sha256':hashlib.sha256(e.content.encode()).hexdigest()}
            if variant=='internal' or variant=='retest' and e.kind=='manual_retest':
                from .workflow_context import structured
                safe,_=redact_object(structured(e.content,12000))
                entry['excerpt']=safe if isinstance(safe,str) else json.dumps(safe,ensure_ascii=False,indent=2)
                entry['truncated']=len(e.content)>12000
            item['evidence'].append(entry)
        if include_images:
            images=db.query(EvidenceAttachment).filter_by(project_id=project.id,finding_id=row.id,attachment_type='screenshot').order_by(EvidenceAttachment.sort_order,EvidenceAttachment.id).all()
            for shot in images:
                try:
                    from PIL import Image
                    raw=render_annotated_image_bytes(shot)
                    with Image.open(io.BytesIO(raw)) as image:
                        output=io.BytesIO();image.save(output,format='PNG');blob=output.getvalue()
                except Exception:raise ValueError(f'截图 #{shot.id} 无法读取，未创建快照。请修复截图或取消包含截图。')
                image_total+=len(blob)
                if image_total>12*1024*1024:raise ValueError('报告截图总量超过 12 MB，请分批交付。')
                item['images'].append({'label':shot.label,'data':base64.b64encode(blob).decode(),'sha256':hashlib.sha256(blob).hexdigest()})
        data['findings'].append(item)
    frozen=encoded(data)
    if len(frozen.encode())>24*1024*1024:raise ValueError('报告内容超过 24 MB，请分批交付。')
    sid=secrets.token_hex(12)
    row=AppPreference(key=f'delivery:{project.id}:{sid}',value_json=encoded({'sha256':hashlib.sha256(frozen.encode()).hexdigest(),'snapshot':data}))
    db.add(row);db.commit()
    return sid

def load_snapshot(db,pid,sid):
    row=db.query(AppPreference).filter_by(key=f'delivery:{pid}:{sid}').first()
    if not row:raise ValueError('交付快照不存在。')
    value=json.loads(row.value_json)
    if hashlib.sha256(encoded(value['snapshot']).encode()).hexdigest()!=value['sha256']:raise ValueError('交付快照完整性校验失败。')
    return value

def report_sections(data):
    yield f"{data['project']} · {VARIANTS[data['variant']]}",f"交付时间：{data['created_at']}\n本报告采用固定快照，后续项目编辑不改变此版本。"
    if data['warnings']:yield '交付前检查','\n'.join(data['warnings'])
    for f in data['findings']:
        body=f"目标：{f['target']}\n风险：{SEVERITY_ZH.get(f['severity'],f['severity'])}\n确认状态：{STATUS_ZH.get(f['finding_state'],f['finding_state'])}\n复测结论：{STATUS_ZH.get(f['verification_state'],f['verification_state'])}\n类型与参数：{f['vuln_type']} / {f['parameter']}\n分类：{f['cwe_id']} · {f['txb02_category']} · {f['owasp_category']}\n\n描述与复现步骤：\n{f['description']}\n\n修复建议：\n{f['recommendation']}"
        for e in f['evidence']:
            body+=f"\n\n证据 #{e['id']} · {e['kind']} · SHA256 {e['sha256']}"
            if 'excerpt' in e:body+='\n'+e['excerpt']+('\n[仅包含前 12000 字符]' if e['truncated'] else '')
        yield f"#{f['id']} {f['title']}",body

def export_snapshot(data,fmt):
    sections=list(report_sections(data))
    if fmt=='md':
        output='\n\n'.join('# '+title+'\n\n'+body for title,body in sections)
        for f in data['findings']:
            for image in f['images']:output+=f"\n\n漏洞 #{f['id']} 截图：{image['label']}\n\n![截图](data:image/png;base64,{image['data']})"
        return output.encode(),'text/markdown; charset=utf-8'
    if fmt=='html':
        esc=html.escape
        content=''.join('<section><h2>'+esc(title)+'</h2><pre>'+esc(body)+'</pre></section>' for title,body in sections)
        for f in data['findings']:
            for image in f['images']:content+=f"<figure><img src='data:image/png;base64,{image['data']}'><figcaption>漏洞 #{f['id']} · {esc(image['label'])}</figcaption></figure>"
        return ("<!doctype html><html lang='zh-CN'><meta charset='utf-8'><meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; img-src data:; style-src 'unsafe-inline'\"><title>交付报告</title><style>body{max-width:960px;margin:32px auto;padding:24px;font:16px/1.7 system-ui}pre{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit}img{max-width:100%}section{border-bottom:1px solid #ddd}</style>"+content+'</html>').encode(),'text/html; charset=utf-8'
    if fmt=='docx':
        from docx import Document
        from docx.shared import Inches,Pt
        document=Document();document.styles['Normal'].font.name='Microsoft YaHei';document.styles['Normal'].font.size=Pt(11)
        for title,body in sections:document.add_heading(title,level=1);document.add_paragraph(body)
        for f in data['findings']:
            for image in f['images']:
                document.add_picture(io.BytesIO(base64.b64decode(image['data'])),width=Inches(6))
                document.add_paragraph(f"漏洞 #{f['id']} · {image['label']}\nSHA256 {image['sha256']}")
        stream=io.BytesIO();document.save(stream)
        return stream.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    raise ValueError('不支持的报告格式。')
