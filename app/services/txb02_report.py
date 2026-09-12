from __future__ import annotations

import io
import json
import re
from pathlib import Path
from collections import Counter
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from sqlalchemy.orm import Session

from ..models import Evidence, EvidenceAttachment, Finding, FindingLifecycle, Project, ProofCapsule, RetestRun
from .evidence_chain import ensure_evidence_integrity, project_evidence_summary
from .evidence_attachments import render_annotated_image_bytes
from .remediation import ensure_project_lifecycles


SENSITIVE_HEADER_RE = re.compile(
    r"(?im)^(authorization|cookie|set-cookie|x-api-key|api-key|x-auth-token|x-access-token)\s*:\s*.*$"
)
STATE_ZH = {
    "candidate": "候选",
    "confirmed": "已确认",
    "false_positive": "误报",
    "open": "待处理",
    "triaged": "已研判",
    "remediation": "整改中",
    "retest_ready": "待复测",
    "resolved": "已修复",
    "accepted_risk": "风险接受",
    "unverified": "未验证",
    "reproduced": "仍可复现",
    "needs_review": "待人工复核",
    "error": "验证失败",
}
SEVERITY_ZH = {"critical": "严重", "high": "高危", "medium": "中危", "low": "低危", "info": "信息"}

SENSITIVE_QUERY = {
    "token", "access_token", "id_token", "refresh_token", "auth", "authorization",
    "code", "ticket", "apikey", "api_key", "key", "session", "sessionid", "sid",
    "jwt", "credential", "secret",
}


def _redact_url(url: str) -> str:
    try:
        parsed = urlparse(url)
        pairs = []
        for key, value in parse_qsl(parsed.query, keep_blank_values=True):
            if key.lower() in SENSITIVE_QUERY:
                pairs.append(f"{key}=••••")
            else:
                pairs.append(f"{key}={value}")
        return urlunparse(parsed._replace(query="&".join(pairs)))
    except Exception:
        return url


def _redact_text(text: str) -> str:
    text = SENSITIVE_HEADER_RE.sub(lambda m: f"{m.group(1)}: ••••", text or "")
    text = re.sub(
        r'(?i)("(?:password|passwd|pwd|token|secret|credential)"\s*:\s*)"[^"]*"',
        r'\1"••••"',
        text,
    )
    # Redact common query secrets even when the URL appears inside a raw HTTP request line.
    for name in sorted(SENSITIVE_QUERY, key=len, reverse=True):
        text = re.sub(
            rf'(?i)([?&]{re.escape(name)}=)([^\s&#]+)',
            lambda m: m.group(1) + "••••",
            text,
        )
    return text[:12000]


def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, size: int = 10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text or "")
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    rpr.append(fonts)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_kv_table(doc: Document, rows: list[tuple[str, str]]):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    for label, value in rows:
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        cells = row.cells
        cells[0].width = Cm(3.6)
        cells[1].width = Cm(13.6)
        _set_cell_shading(cells[0], "E9EEF3")
        _set_cell_text(cells[0], label, bold=True, size=10)
        _set_cell_text(cells[1], value or "未记录", size=10)
    return table


def _evidence_sections(db: Session, finding: Finding) -> dict:
    rows = (
        db.query(Evidence)
        .filter(Evidence.finding_id == finding.id)
        .order_by(Evidence.id.asc())
        .all()
    )
    result = {
        "normal_request": "",
        "normal_response": "",
        "poc_request": "",
        "poc_response": "",
        "detail": [],
        "digests": [],
    }
    for evidence in rows:
        ensure_evidence_integrity(db, evidence)
        kind = (evidence.kind or "").lower()
        content = _redact_text(evidence.content or "")
        digest = evidence.content_sha256 or ""
        short_digest = (digest[:16] + "…" + digest[-8:]) if len(digest) == 64 else digest
        result["digests"].append(f"Evidence #{evidence.id} · {evidence.kind} · SHA256 {short_digest}")
        if not result["normal_request"] and ("baseline_request" in kind or "normal_request" in kind):
            result["normal_request"] = content
        elif not result["normal_response"] and ("baseline_response" in kind or "normal_response" in kind):
            result["normal_response"] = content
        elif not result["poc_request"] and ("poc_request" in kind or "request_replay" in kind):
            result["poc_request"] = content
        elif not result["poc_response"] and ("poc_response" in kind or "authorization_candidate" in kind):
            result["poc_response"] = content
        elif kind == "screenshot_evidence":
            continue
        elif len(result["detail"]) < 4:
            result["detail"].append(f"[{evidence.kind}] {content[:1800]}")
    return result


def _difficulty(finding: Finding) -> str:
    if finding.source == "authorization_testing":
        return "中"
    if finding.severity == "high":
        return "中"
    if finding.severity == "medium":
        return "中"
    return "低"


def generate_txb02_docx(db: Session, project: Project) -> bytes:
    ensure_project_lifecycles(db, project.id)
    evidence_summary = project_evidence_summary(db, project.id)
    findings = (
        db.query(Finding)
        .filter(Finding.project_id == project.id)
        .order_by(Finding.id.asc())
        .all()
    )
    lifecycle_map = {
        x.finding_id: x
        for x in db.query(FindingLifecycle).filter(FindingLifecycle.project_id == project.id).all()
    }
    capsules = {
        x.finding_id: x
        for x in db.query(ProofCapsule).filter(ProofCapsule.project_id == project.id).all()
    }

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Microsoft YaHei"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Microsoft YaHei"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Microsoft YaHei"
    styles["Heading 2"].font.size = Pt(13)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("渗透测试报告")
    run.bold = True
    run.font.size = Pt(22)
    run.font.name = "Microsoft YaHei"

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(project.name).bold = True

    meta = [
        ("客户名称", project.client_name or "未填写"),
        ("项目名称", project.name),
        ("测试环境", project.environment),
        ("项目状态", project.status),
        ("测试时间", f"{project.start_date or '—'} 至 {project.end_date or '—'}"),
        ("授权范围", project.scope_text),
        ("授权说明", project.authorization_note or "未填写"),
    ]
    _add_kv_table(doc, meta)

    doc.add_paragraph()
    h = doc.add_heading("一、测试结果概览", level=1)
    sev = Counter((x.severity or "info").lower() for x in findings)
    _add_kv_table(doc, [
        ("漏洞总数", str(len(findings))),
        ("高危", str(sev.get("high", 0))),
        ("中危", str(sev.get("medium", 0))),
        ("低危", str(sev.get("low", 0))),
        ("信息", str(sev.get("info", 0))),
        ("Evidence", str(evidence_summary["count"])),
        ("Evidence SHA256 完整率", f"{evidence_summary['integrity_ready']}/{evidence_summary['count']}"),
    ])

    doc.add_heading("二、漏洞详情", level=1)
    if not findings:
        doc.add_paragraph("本次项目当前未记录漏洞。")

    for idx, finding in enumerate(findings, start=1):
        lifecycle = lifecycle_map.get(finding.id)
        capsule = capsules.get(finding.id)
        latest_retest = None
        if capsule:
            latest_retest = (
                db.query(RetestRun)
                .filter(RetestRun.proof_capsule_id == capsule.id)
                .order_by(RetestRun.id.desc())
                .first()
            )
        evidence = _evidence_sections(db, finding)
        attachments = (
            db.query(EvidenceAttachment)
            .filter(EvidenceAttachment.finding_id == finding.id, EvidenceAttachment.attachment_type == "screenshot")
            .order_by(EvidenceAttachment.sort_order.asc(), EvidenceAttachment.id.asc())
            .all()
        )
        detail_text = finding.description
        if evidence["detail"]:
            detail_text += "\n\n证据摘要：\n" + "\n".join(evidence["detail"])

        doc.add_heading(f"{idx}. {finding.title}", level=2)
        _add_kv_table(doc, [
            ("漏洞标题", finding.title),
            ("漏洞URL", _redact_url(finding.target)),
            ("漏洞参数", finding.parameter or "不涉及 / 未记录"),
            ("漏洞类型", finding.txb02_category or finding.vuln_type or "待人工分类"),
            ("CWE", finding.cwe_id or "未映射"),
            ("OWASP", finding.owasp_category or "未映射"),
            ("利用难度", _difficulty(finding)),
            ("漏洞等级", SEVERITY_ZH.get(finding.severity, finding.severity.upper())),
            ("漏洞状态", STATE_ZH.get(finding.finding_state, finding.finding_state)),
            ("整改状态", STATE_ZH.get(lifecycle.status if lifecycle else "open", lifecycle.status if lifecycle else "open")),
            ("验证状态", STATE_ZH.get(finding.verification_state, finding.verification_state)),
            ("Evidence 摘要", "; ".join(evidence["digests"]) if evidence["digests"] else "未记录"),
            ("漏洞详述", detail_text),
            ("正常请求", evidence["normal_request"] or "未记录"),
            ("正常响应", evidence["normal_response"] or "未记录"),
            ("漏洞POC请求", evidence["poc_request"] or "未记录"),
            ("漏洞POC响应", evidence["poc_response"] or "未记录"),
            ("验证截图", f"{len(attachments)} 张" if attachments else "未记录"),
            ("修复建议", finding.recommendation or "未填写"),
            ("复测结果", latest_retest.detail if latest_retest else "未复测 / 未记录"),
        ])
        if attachments:
            shot_heading = doc.add_paragraph(); shot_run = shot_heading.add_run("验证截图"); shot_run.bold = True; shot_run.font.size = Pt(11); shot_run.font.name = "Microsoft YaHei"
            for shot_index, attachment in enumerate(attachments, start=1):
                image_path = Path(attachment.file_path)
                if not image_path.exists(): continue
                try:
                    try:
                        rendered = render_annotated_image_bytes(attachment)
                        pic = doc.add_picture(io.BytesIO(rendered), width=Inches(6.1))
                    except Exception:
                        # Annotation rendering is an enhancement. A renderer
                        # incompatibility must not regress V1.4 screenshot delivery.
                        pic = doc.add_picture(str(image_path), width=Inches(6.1))
                    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption = doc.add_paragraph(); caption.alignment = WD_ALIGN_PARAGRAPH.CENTER; cap = caption.add_run(f"图 {idx}-{shot_index} {attachment.label or '漏洞验证截图'}"); cap.font.size = Pt(9.5); cap.font.name = "Microsoft YaHei"
                    digest = doc.add_paragraph(); digest.alignment = WD_ALIGN_PARAGRAPH.CENTER; digest_run = digest.add_run(f"SHA256 {attachment.file_sha256}"); digest_run.font.size = Pt(8); digest_run.font.name = "Cascadia Code"
                except Exception:
                    doc.add_paragraph(f"[截图无法嵌入] {attachment.label} · SHA256 {attachment.file_sha256}")
        if idx < len(findings):
            doc.add_paragraph()

    # The report deliberately omits a trailing boilerplate page. Security and
    # redaction semantics are documented in the Workspace and delivery metadata.

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()
